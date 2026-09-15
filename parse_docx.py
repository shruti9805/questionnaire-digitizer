"""Parse a phase's bilingual questionnaire .docx template into a schema.

The template has a fixed structure (see PRODUCT.md): table[0] is the
demographic info table (label cell / Hindi-label cell per row, some rows
carry checkbox options separated by the "☐" glyph); table[1] is the
Likert item table (header row, then one row per item: SrNo | Code |
Statement (English/Hindi) | 1 | 2 | 3 | 4 | 5). Verified against the real
Student_Survey_Bilingual.docx this session: each statement cell holds one
paragraph with a line-break (not a paragraph break) between the English and
Hindi text, so splitting cell.text on the first "\n" reliably separates them.
"""
from __future__ import annotations

from dataclasses import dataclass

CHECKBOX_GLYPH = "☐"  # ☐


@dataclass
class ParsedItem:
    position: int
    code: str
    statement_en: str
    statement_hi: str | None


@dataclass
class ParsedDemoField:
    position: int
    label_en: str
    label_hi: str | None
    field_type: str  # "text" | "single_choice"
    options: list[str] | None


class SchemaParseError(Exception):
    pass


def _split_en_hi(cell_text: str) -> tuple[str, str | None]:
    parts = cell_text.split("\n", 1)
    en = parts[0].strip()
    hi = parts[1].strip() if len(parts) > 1 else None
    return en, (hi or None)


def _parse_demo_cell(cell_text: str) -> tuple[str, str, list[str] | None]:
    """Returns (label, field_type, options) for one language's cell text."""
    lines = [ln for ln in cell_text.split("\n")]
    label = lines[0].strip()
    remainder = "\n".join(lines[1:])
    if CHECKBOX_GLYPH in remainder:
        raw_opts = remainder.split(CHECKBOX_GLYPH)
        options = [opt.strip(" ,\n\t") for opt in raw_opts]
        options = [opt for opt in options if opt]
        return label, "single_choice", options
    return label, "text", None


def parse_phase_docx(path: str) -> tuple[list[ParsedItem], list[ParsedDemoField]]:
    import docx  # local import: keep this module importable without the dep for callers that don't need it

    document = docx.Document(path)
    if len(document.tables) < 2:
        raise SchemaParseError(
            f"Expected at least 2 tables (demographics, Likert items); found {len(document.tables)}. "
            "This doesn't look like the expected phase template."
        )

    demo_table, item_table = document.tables[0], document.tables[1]

    demo_fields: list[ParsedDemoField] = []
    for i, row in enumerate(demo_table.rows):
        if len(row.cells) < 2:
            raise SchemaParseError(f"Demographics table row {i} has fewer than 2 columns (EN/HI).")
        en_label, field_type, options = _parse_demo_cell(row.cells[0].text)
        hi_label, _, _ = _parse_demo_cell(row.cells[1].text)
        if not en_label:
            raise SchemaParseError(f"Demographics table row {i} has an empty English label.")
        demo_fields.append(
            ParsedDemoField(
                position=i,
                label_en=en_label,
                label_hi=hi_label or None,
                field_type=field_type,
                options=options,
            )
        )

    header_cells = [c.text.strip() for c in item_table.rows[0].cells]
    expected_header = ["SrNo", "Code", "Statement (English / Hindi)", "1", "2", "3", "4", "5"]
    if header_cells != expected_header:
        raise SchemaParseError(
            f"Likert table header doesn't match the expected template.\n"
            f"Expected: {expected_header}\nFound:    {header_cells}"
        )

    items: list[ParsedItem] = []
    for i, row in enumerate(item_table.rows[1:], start=1):
        if len(row.cells) < 3:
            raise SchemaParseError(f"Likert table row {i} has fewer than 3 columns.")
        code = row.cells[1].text.strip()
        if not code:
            raise SchemaParseError(f"Likert table row {i} has an empty Code cell.")
        statement_en, statement_hi = _split_en_hi(row.cells[2].text)
        if not statement_en:
            raise SchemaParseError(f"Likert table row {i} (code={code}) has an empty statement.")
        items.append(
            ParsedItem(position=i, code=code, statement_en=statement_en, statement_hi=statement_hi)
        )

    if not items:
        raise SchemaParseError("No Likert items found in the template.")

    return items, demo_fields
