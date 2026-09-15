"""Export a reviewed response to an Excel workbook, matching the structure
hand-built earlier this session (Demographics / Likert_Responses /
Wide_Format / Extraction_Notes sheets) but driven from the DB schema instead
of a hardcoded item list, so it works for any phase's questionnaire.
"""
from __future__ import annotations

from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from db import get_demo_fields, get_items, get_response_demo_values, get_response_items, list_responses

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
SCALE_LABEL = {1: "Strongly Disagree", 2: "Disagree", 3: "Neither Agree nor Disagree",
               4: "Agree", 5: "Strongly Agree", None: "(blank)"}


def _header_row(ws, row_idx, headers):
    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=row_idx, column=j, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


def build_workbook(phase_id: int, phase_name: str, source_pdf_filename: str, response_id: int) -> openpyxl.Workbook:
    demo_fields = get_demo_fields(phase_id)
    demo_values = get_response_demo_values(response_id)
    items = get_items(phase_id)
    response_items = {ri["item_id"]: ri for ri in get_response_items(response_id)}

    wb = openpyxl.Workbook()

    # ---------- Demographics ----------
    ws1 = wb.active
    ws1.title = "Demographics"
    _header_row(ws1, 1, ["Field", "Value"])
    for i, df in enumerate(demo_fields, start=2):
        ws1.cell(row=i, column=1, value=df["label_en"]).border = BORDER
        ws1.cell(row=i, column=2, value=demo_values.get(df["id"], "")).border = BORDER
    ws1.column_dimensions["A"].width = 45
    ws1.column_dimensions["B"].width = 55
    for row in ws1.iter_rows(min_row=1, max_row=len(demo_fields) + 1, min_col=1, max_col=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # ---------- Likert_Responses (long format) ----------
    ws2 = wb.create_sheet("Likert_Responses")
    _header_row(ws2, 1, ["Sr.No", "Code", "Statement (English)", "Response (1-5)", "Confidence", "Note"])
    for i, it in enumerate(items, start=2):
        ri = response_items.get(it["id"])
        value = ri["value"] if ri else None
        confidence = ri["confidence"] if ri else "missing"
        note = ri["note"] if ri else "no detection recorded for this item"
        ws2.cell(row=i, column=1, value=it["position"]).border = BORDER
        ws2.cell(row=i, column=2, value=it["code"]).border = BORDER
        ws2.cell(row=i, column=3, value=it["statement_en"]).border = BORDER
        vcell = ws2.cell(row=i, column=4, value=value)
        vcell.border = BORDER
        vcell.alignment = Alignment(horizontal="center")
        ws2.cell(row=i, column=5, value=confidence).border = BORDER
        ws2.cell(row=i, column=6, value=note).border = BORDER
    ws2.column_dimensions["A"].width = 8
    ws2.column_dimensions["B"].width = 10
    ws2.column_dimensions["C"].width = 80
    ws2.column_dimensions["D"].width = 14
    ws2.column_dimensions["E"].width = 12
    ws2.column_dimensions["F"].width = 45
    for row in ws2.iter_rows(min_row=1, max_row=len(items) + 1):
        for cell in row:
            if cell.column in (3, 6):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.freeze_panes = "A2"

    # ---------- Wide_Format (one row per respondent) ----------
    ws3 = wb.create_sheet("Wide_Format")
    demo_headers = [df["label_en"] for df in demo_fields]
    item_headers = [it["code"] for it in items]
    headers = ["Respondent_File"] + demo_headers + item_headers
    _header_row(ws3, 1, headers)
    row_values = [source_pdf_filename]
    row_values += [demo_values.get(df["id"], "") for df in demo_fields]
    for it in items:
        ri = response_items.get(it["id"])
        row_values.append(ri["value"] if ri else None)
    for j, val in enumerate(row_values, start=1):
        ws3.cell(row=2, column=j, value=val).border = BORDER
    for j in range(1, len(headers) + 1):
        ws3.column_dimensions[get_column_letter(j)].width = 14
    ws3.freeze_panes = "B2"

    # ---------- Extraction_Notes ----------
    ws4 = wb.create_sheet("Extraction_Notes")
    flagged = [ri for ri in response_items.values() if ri["confidence"] != "ok"]
    notes = [
        ("Phase", f"{phase_name} (id {phase_id})"),
        ("Source PDF", source_pdf_filename),
        ("Method", "Pages rendered at 300 DPI (pdftoppm). Likert checkbox values (1-5) detected by: "
                   "(1) locating the 5 answer-column gridlines via longest-contiguous-dark-run projection, "
                   "(2) isolating handwritten ink via a per-document Otsu threshold on HSV saturation "
                   "(pen-color-agnostic, not hardcoded to any specific ink color), (3) clustering ink "
                   "pixels into one mark per answered row and binning by column, (4) reconciling the total "
                   "detected count against this phase's known item count before trusting any alignment."),
        ("Demographic fields", "Not automatically read in this version (no handwriting/OCR recognition) - "
                                "entered manually by the researcher while viewing the scanned page."),
        ("Flagged for review", f"{len(flagged)} of {len(items)} items were flagged as low-confidence or "
                                f"conflicting during automated detection (see below) and required a human "
                                f"confirm/correct pass before this export."),
    ]
    ws4.column_dimensions["A"].width = 28
    ws4.column_dimensions["B"].width = 110
    row = 1
    for k, v in notes:
        ws4.cell(row=row, column=1, value=k).font = Font(bold=True)
        ws4.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws4.cell(row=row, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    row += 1
    if flagged:
        ws4.cell(row=row, column=1, value="Flagged items (resolved before export)").font = Font(bold=True)
        row += 1
        _header_row(ws4, row, ["Code", "Value", "Confidence", "Note"])
        row += 1
        code_by_item_id = {it["id"]: it["code"] for it in items}
        for ri in sorted(flagged, key=lambda r: r["item_id"]):
            ws4.cell(row=row, column=1, value=code_by_item_id.get(ri["item_id"], "?")).border = BORDER
            ws4.cell(row=row, column=2, value=ri["value"]).border = BORDER
            ws4.cell(row=row, column=3, value=ri["confidence"]).border = BORDER
            note_cell = ws4.cell(row=row, column=4, value=ri["note"])
            note_cell.border = BORDER
            note_cell.alignment = Alignment(wrap_text=True, vertical="top")
            row += 1

    return wb


def export_response_to_bytes(phase_id: int, phase_name: str, source_pdf_filename: str, response_id: int) -> bytes:
    wb = build_workbook(phase_id, phase_name, source_pdf_filename, response_id)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_batch_workbook(phase_id: int, phase_name: str, batch_id: int) -> openpyxl.Workbook:
    """Same four-sheet shape as build_workbook, generalized to N respondents:
    Demographics and Wide_Format each get one row per respondent instead of
    one row per field; Likert_Responses gains a Respondent column."""
    demo_fields = get_demo_fields(phase_id)
    items = get_items(phase_id)
    responses = list_responses(batch_id)

    wb = openpyxl.Workbook()

    def respondent_name(r):
        return r["respondent_label"] or r["source_pdf_filename"]

    # ---------- Demographics (one row per respondent) ----------
    ws1 = wb.active
    ws1.title = "Demographics"
    _header_row(ws1, 1, ["Respondent"] + [df["label_en"] for df in demo_fields])
    for i, r in enumerate(responses, start=2):
        demo_values = get_response_demo_values(r["id"])
        ws1.cell(row=i, column=1, value=respondent_name(r)).border = BORDER
        for j, df in enumerate(demo_fields, start=2):
            ws1.cell(row=i, column=j, value=demo_values.get(df["id"], "")).border = BORDER
    ws1.column_dimensions["A"].width = 30
    for j in range(2, len(demo_fields) + 2):
        ws1.column_dimensions[get_column_letter(j)].width = 30
    ws1.freeze_panes = "B2"

    # ---------- Likert_Responses (long format, all respondents) ----------
    ws2 = wb.create_sheet("Likert_Responses")
    _header_row(ws2, 1, ["Respondent", "Sr.No", "Code", "Statement (English)", "Response (1-5)", "Confidence", "Note"])
    row = 2
    for r in responses:
        response_items = {ri["item_id"]: ri for ri in get_response_items(r["id"])}
        for it in items:
            ri = response_items.get(it["id"])
            value = ri["value"] if ri else None
            confidence = ri["confidence"] if ri else "missing"
            note = ri["note"] if ri else "no detection recorded for this item"
            ws2.cell(row=row, column=1, value=respondent_name(r)).border = BORDER
            ws2.cell(row=row, column=2, value=it["position"]).border = BORDER
            ws2.cell(row=row, column=3, value=it["code"]).border = BORDER
            ws2.cell(row=row, column=4, value=it["statement_en"]).border = BORDER
            vcell = ws2.cell(row=row, column=5, value=value)
            vcell.border = BORDER
            vcell.alignment = Alignment(horizontal="center")
            ws2.cell(row=row, column=6, value=confidence).border = BORDER
            ws2.cell(row=row, column=7, value=note).border = BORDER
            row += 1
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 8
    ws2.column_dimensions["C"].width = 10
    ws2.column_dimensions["D"].width = 70
    ws2.column_dimensions["E"].width = 14
    ws2.column_dimensions["F"].width = 12
    ws2.column_dimensions["G"].width = 45
    ws2.freeze_panes = "A2"

    # ---------- Wide_Format (one row per respondent) ----------
    ws3 = wb.create_sheet("Wide_Format")
    demo_headers = [df["label_en"] for df in demo_fields]
    item_headers = [it["code"] for it in items]
    headers = ["Respondent_File"] + demo_headers + item_headers
    _header_row(ws3, 1, headers)
    for i, r in enumerate(responses, start=2):
        demo_values = get_response_demo_values(r["id"])
        response_items = {ri["item_id"]: ri for ri in get_response_items(r["id"])}
        row_values = [respondent_name(r)]
        row_values += [demo_values.get(df["id"], "") for df in demo_fields]
        for it in items:
            ri = response_items.get(it["id"])
            row_values.append(ri["value"] if ri else None)
        for j, val in enumerate(row_values, start=1):
            ws3.cell(row=i, column=j, value=val).border = BORDER
    for j in range(1, len(headers) + 1):
        ws3.column_dimensions[get_column_letter(j)].width = 14
    ws3.freeze_panes = "B2"

    # ---------- Extraction_Notes ----------
    ws4 = wb.create_sheet("Extraction_Notes")
    notes = [
        ("Phase", f"{phase_name} (id {phase_id})"),
        ("Batch", f"batch id {batch_id}, {len(responses)} respondent(s)"),
        ("Method", "Pages rendered at 300 DPI (pdftoppm). Likert checkbox values (1-5) detected by: "
                   "(1) locating the 5 answer-column gridlines via longest-contiguous-dark-run projection, "
                   "(2) isolating handwritten ink via a per-document Otsu threshold on HSV saturation "
                   "(pen-color-agnostic, not hardcoded to any specific ink color), (3) clustering ink "
                   "pixels into one mark per answered row and binning by column, (4) reconciling the total "
                   "detected count against this phase's known item count before trusting any alignment."),
        ("Demographic fields", "Not automatically read in this version (no handwriting/OCR recognition) - "
                                "entered manually by the researcher while viewing each scanned response."),
    ]
    ws4.column_dimensions["A"].width = 28
    ws4.column_dimensions["B"].width = 110
    row = 1
    for k, v in notes:
        ws4.cell(row=row, column=1, value=k).font = Font(bold=True)
        ws4.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws4.cell(row=row, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    row += 1
    ws4.cell(row=row, column=1, value="Per-respondent flag summary").font = Font(bold=True)
    row += 1
    _header_row(ws4, row, ["Respondent", "Items", "Flagged"])
    row += 1
    for r in responses:
        response_items = get_response_items(r["id"])
        n_flagged = sum(1 for ri in response_items if ri["confidence"] != "ok")
        ws4.cell(row=row, column=1, value=respondent_name(r)).border = BORDER
        ws4.cell(row=row, column=2, value=len(response_items)).border = BORDER
        ws4.cell(row=row, column=3, value=n_flagged).border = BORDER
        row += 1

    return wb


def export_batch_to_bytes(phase_id: int, phase_name: str, batch_id: int) -> bytes:
    wb = build_batch_workbook(phase_id, phase_name, batch_id)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
